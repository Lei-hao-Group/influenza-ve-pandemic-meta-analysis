#!/usr/bin/env python3
"""Create the Newcastle–Ottawa Scale bias summary figure."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Patch
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
INPUT = (
    ROOT
    / "Submission_Code"
    / "Data"
    / "Supplementary Data_Characteristics and Quality score.xlsx"
)
OUTPUT_DIR = ROOT / "Results" / "BiasAssessment"
FONT_SIZE = 11

for font_file in (
    "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/Arial_Italic.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold_Italic.ttf",
):
    if Path(font_file).exists():
        font_manager.fontManager.addfont(font_file)

CATEGORIES = (
    "Maximum stars awarded",
    "One of two stars awarded",
    "No star awarded",
)
COLORS = {
    CATEGORIES[0]: "#90D4C9",
    CATEGORIES[1]: "#FFCB8F",
    CATEGORIES[2]: "#E3DED3",
}

DESIGNS = (
    {
        "sheet": "NOS case-control",
        "name": "Case-control studies",
        "items": (
            ("Selection", "S1  Case definition"),
            ("Selection", "S2  Representative cases"),
            ("Selection", "S3  Control selection"),
            ("Selection", "S4  Control definition"),
            ("Comparability", "C1  Comparability"),
            ("Exposure", "E1  Exposure ascertainment"),
            ("Exposure", "E2  Same ascertainment method"),
            ("Exposure", "E3  Non-response rate"),
        ),
    },
    {
        "sheet": "NOS cohort",
        "name": "Cohort studies",
        "items": (
            ("Selection", "S1  Exposed cohort"),
            ("Selection", "S2  Non-exposed cohort"),
            ("Selection", "S3  Exposure ascertainment"),
            ("Selection", "S4  Outcome absent at baseline"),
            ("Comparability", "C1  Comparability"),
            ("Outcome", "O1  Outcome assessment"),
            ("Outcome", "O2  Follow-up duration"),
            ("Outcome", "O3  Follow-up completeness"),
        ),
    },
)


def read_design_rows(ws):
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[1] is None:
            continue
        ratings = row[2:10]
        if len(ratings) != 8 or any(v is None for v in ratings):
            raise ValueError(f"Missing NOS rating in {ws.title}: {row[:2]}")
        if any(v not in (0, 1, 2) for v in ratings):
            raise ValueError(f"Unexpected NOS rating in {ws.title}: {row[:2]} {ratings}")
        if any(v not in (0, 1) for i, v in enumerate(ratings) if i != 4):
            raise ValueError(f"Only comparability may have two stars: {row[:2]} {ratings}")
        rows.append(ratings)
    return rows


def classify(value, comparability=False):
    if comparability:
        return CATEGORIES[{2: 0, 1: 1, 0: 2}[value]]
    return CATEGORIES[0] if value == 1 else CATEGORIES[2]


def summarize(workbook):
    summaries = []
    all_scores = []
    for design in DESIGNS:
        rows = read_design_rows(workbook[design["sheet"]])
        scores = [sum(row) for row in rows]
        all_scores.extend(scores)
        item_summaries = []
        for item_index, (domain, label) in enumerate(design["items"]):
            counts = Counter(
                classify(row[item_index], comparability=(item_index == 4))
                for row in rows
            )
            item_summaries.append(
                {
                    "domain": domain,
                    "label": label,
                    "counts": {category: counts[category] for category in CATEGORIES},
                    "percents": {
                        category: counts[category] * 100.0 / len(rows)
                        for category in CATEGORIES
                    },
                }
            )
        summaries.append(
            {
                **design,
                "n": len(rows),
                "items_summary": item_summaries,
                "score_counts": Counter(scores),
            }
        )
    return summaries, Counter(all_scores)


def export_source_data(summaries, overall_score_counts):
    path = OUTPUT_DIR / "Summary_of_bias_assessment_source_data.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["design", "n_studies", "domain", "item", "rating", "count", "percent"]
        )
        for design in summaries:
            for item in design["items_summary"]:
                for category in CATEGORIES:
                    writer.writerow(
                        [
                            design["name"],
                            design["n"],
                            item["domain"],
                            item["label"],
                            category,
                            item["counts"][category],
                            f'{item["percents"][category]:.4f}',
                        ]
                    )
        n_all = sum(overall_score_counts.values())
        for score in range(5, 10):
            count = overall_score_counts[score]
            writer.writerow(
                [
                    "All studies",
                    n_all,
                    "Overall",
                    "Overall NOS score",
                    f"{score} points",
                    count,
                    f"{count * 100.0 / n_all:.4f}",
                ]
            )
    return path


def plot_summary(summaries, overall_score_counts):
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": FONT_SIZE,
            "axes.titlesize": FONT_SIZE,
            "axes.titleweight": "normal",
            "axes.linewidth": 0.5,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none",
        }
    )

    # Nature double-column dimensions: 183 mm wide and below the 170 mm height limit.
    fig, axes = plt.subplots(
        3,
        1,
        figsize=(183 / 25.4, 166 / 25.4),
        gridspec_kw={"height_ratios": (1, 1, 0.48)},
        constrained_layout=False,
    )
    fig.patch.set_facecolor("white")

    for ax, design in zip(axes[:2], summaries):
        items = design["items_summary"]
        y = list(range(len(items)))
        left = [0.0] * len(items)

        for category in CATEGORIES:
            widths = [item["percents"][category] for item in items]
            bars = ax.barh(
                y,
                widths,
                left=left,
                height=0.62,
                color=COLORS[category],
                edgecolor="white",
                linewidth=0.45,
                zorder=3,
            )
            for bar, width in zip(bars, widths):
                if width >= 7.5:
                    txt_color = "#111111"
                    ax.text(
                        bar.get_x() + bar.get_width() / 2,
                        bar.get_y() + bar.get_height() / 2,
                        f"{width:.0f}%",
                        ha="center",
                        va="center",
                        color=txt_color,
                        fontsize=FONT_SIZE,
                        fontweight="normal",
                        zorder=4,
                    )
            left = [a + b for a, b in zip(left, widths)]

        tick_labels = [item["label"] for item in items]
        ax.set_yticks(y, tick_labels)
        ax.invert_yaxis()
        ax.set_xlim(0, 100)
        ax.set_xticks(range(0, 101, 25))
        ax.set_xticklabels([str(v) for v in range(0, 101, 25)])
        ax.set_xlabel("Studies (%)", labelpad=2, fontsize=FONT_SIZE)
        ax.xaxis.label.set_fontweight("bold")
        ax.set_title(
            design["name"],
            loc="center",
            fontsize=FONT_SIZE,
            fontweight="bold",
            pad=4,
        )
        ax.xaxis.grid(True, color="#D6D6D6", linewidth=0.4, zorder=0)
        ax.yaxis.grid(False)
        ax.tick_params(axis="y", length=0, pad=4, labelsize=FONT_SIZE)
        ax.tick_params(
            axis="x",
            width=0.5,
            length=2.2,
            pad=1.5,
            colors="#333333",
            labelsize=FONT_SIZE,
        )
        for spine in ("top", "right", "left"):
            ax.spines[spine].set_visible(False)
        ax.spines["bottom"].set_color("#777777")
        ax.spines["bottom"].set_linewidth(0.5)

        # Thin separators emphasize the NOS domains without adding visual clutter.
        domains = [item["domain"] for item in items]
        for idx in range(len(domains) - 1):
            if domains[idx] != domains[idx + 1]:
                ax.axhline(idx + 0.5, color="#9A9A9A", linewidth=0.45, zorder=1)

    score_ax = axes[2]
    scores = list(range(5, 10))
    n_all = sum(overall_score_counts.values())
    counts = [overall_score_counts[score] for score in scores]
    percents = [count * 100.0 / n_all for count in counts]
    score_bars = score_ax.bar(
        scores,
        percents,
        width=0.66,
        color=COLORS[CATEGORIES[0]],
        edgecolor="white",
        linewidth=0.45,
        zorder=3,
    )
    for bar, count, percent in zip(score_bars, counts, percents):
        score_ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + 0.8,
            f"{count} ({percent:.1f}%)",
            ha="center",
            va="bottom",
            fontsize=FONT_SIZE,
            color="#222222",
        )
    score_ax.set_xlim(4.5, 9.5)
    score_ax.set_ylim(0, max(percents) + 9)
    score_ax.axvline(6.5, color="#777777", linewidth=0.7, linestyle="--", zorder=2)
    score_ax.set_xticks(scores)
    score_ax.set_xlabel("NOS score (0–9)", labelpad=2, fontsize=FONT_SIZE)
    score_ax.xaxis.label.set_fontweight("bold")
    score_ax.set_ylabel("Studies (%)", labelpad=5, fontsize=FONT_SIZE)
    score_ax.yaxis.label.set_fontweight("bold")
    score_ax.set_title(
        f"Overall NOS score (all studies, n={n_all})",
        loc="center",
        fontsize=FONT_SIZE,
        fontweight="bold",
        pad=4,
    )
    score_ax.yaxis.grid(True, color="#D6D6D6", linewidth=0.4, zorder=0)
    score_ax.xaxis.grid(False)
    score_ax.tick_params(
        axis="both",
        width=0.5,
        length=2.2,
        pad=1.5,
        colors="#333333",
        labelsize=FONT_SIZE,
    )
    for spine in ("top", "right"):
        score_ax.spines[spine].set_visible(False)
    for spine in ("bottom", "left"):
        score_ax.spines[spine].set_color("#777777")
        score_ax.spines[spine].set_linewidth(0.5)

    short_labels = ("Maximum stars", "One of two stars", "No star")
    legend_handles = [
        Patch(facecolor=COLORS[c], edgecolor="none", label=label)
        for c, label in zip(CATEGORIES, short_labels)
    ]
    fig.legend(
        handles=legend_handles,
        loc="upper right",
        bbox_to_anchor=(0.985, 0.985),
        ncol=3,
        frameon=False,
        columnspacing=1.2,
        handlelength=1.25,
        handleheight=0.8,
        borderaxespad=0,
        fontsize=FONT_SIZE,
    )
    fig.subplots_adjust(left=0.34, right=0.975, top=0.91, bottom=0.08, hspace=0.42)

    stem = OUTPUT_DIR / "Summary_of_bias_assessment"
    png_path = stem.with_suffix(".png")
    tiff_path = stem.with_suffix(".tiff")
    fig.savefig(png_path, dpi=300, facecolor="white")
    fig.savefig(tiff_path, dpi=600, facecolor="white", pil_kwargs={"compression": "tiff_lzw"})
    fig.savefig(stem.with_suffix(".pdf"), facecolor="white")
    fig.savefig(stem.with_suffix(".svg"), facecolor="white")
    plt.close(fig)

    # Journal raster deliverables are saved as opaque RGB rather than RGBA.
    for image_path, dpi in ((png_path, 300), (tiff_path, 600)):
        with Image.open(image_path) as raster:
            rgb = raster.convert("RGB").copy()
        temporary_path = image_path.with_name(f"{image_path.stem}.rgb{image_path.suffix}")
        if image_path.suffix == ".tiff":
            rgb.save(str(temporary_path), dpi=(dpi, dpi), compression="tiff_lzw")
        else:
            rgb.save(str(temporary_path), dpi=(dpi, dpi))
        temporary_path.replace(image_path)
    return stem


def export_caption(summaries, overall_score_counts):
    case_n, cohort_n = (d["n"] for d in summaries)
    n_all = sum(overall_score_counts.values())
    n_ge7 = sum(count for score, count in overall_score_counts.items() if score >= 7)
    caption = (
        "Summary of Newcastle–Ottawa Scale item ratings. Item-level NOS star awards for "
        f"{case_n} case-control studies and {cohort_n} cohort studies. Bars show the "
        "percentage of studies awarded the maximum number of stars, one of two possible stars, "
        "or no star for each item; S, C, E and O denote Selection, Comparability, Exposure and "
        "Outcome, respectively. Selection and Exposure/Outcome items permit a maximum of one star "
        "each, whereas the Comparability item permits a maximum of two stars; consequently, the "
        "one-of-two-stars category applies only to Comparability. These are native NOS star awards "
        "and are not presented as low, moderate or high risk-of-bias judgements. "
        f"The lower panel shows the distribution of summed NOS scores among all {n_all} studies; "
        f"{n_ge7} ({n_ge7 * 100.0 / n_all:.1f}%) scored at least 7 of 9. Bar labels give the "
        "number and percentage of studies. The dashed line marks the NOS score of at least 7 used "
        "for the prespecified high-quality sensitivity analysis; it is not a native NOS risk-of-bias cutoff."
    )
    path = OUTPUT_DIR / "Summary_of_bias_assessment_caption.txt"
    path.write_text(caption + "\n", encoding="utf-8")
    return path


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(INPUT, data_only=True, read_only=True)
    summaries, overall_score_counts = summarize(workbook)
    if [d["n"] for d in summaries] != [323, 26]:
        raise ValueError(f'Unexpected study counts: {[d["n"] for d in summaries]}')
    expected_score_counts = {5: 8, 6: 38, 7: 87, 8: 150, 9: 66}
    if dict(sorted(overall_score_counts.items())) != expected_score_counts:
        raise ValueError(f"Unexpected overall NOS score counts: {overall_score_counts}")
    source_path = export_source_data(summaries, overall_score_counts)
    figure_stem = plot_summary(summaries, overall_score_counts)
    caption_path = export_caption(summaries, overall_score_counts)
    print(f"Created {figure_stem}.png/.tiff/.pdf/.svg")
    print(f"Created {source_path}")
    print(f"Created {caption_path}")


if __name__ == "__main__":
    main()
